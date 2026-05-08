"""workspace 类 skills：CSV、JSON、git 等（依赖可选 openpyxl/python-docx）。"""

from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path

from webapi.skill_runner import load_skills_registry, run_skill_call


def _root() -> Path:
    return Path(__file__).resolve().parents[1]


class WorkspaceSkillsTest(unittest.TestCase):
    def setUp(self) -> None:
        self.funcs, _ = load_skills_registry()

    def test_csv_tsv_summary(self) -> None:
        with tempfile.NamedTemporaryFile("w", suffix=".csv", delete=False, encoding="utf-8") as f:
            f.write("a,b\n1,2\n3,4\n")
            p = f.name
        try:
            ret = run_skill_call(f"csv_tsv_summary({p!r})", self.funcs)
            self.assertEqual(ret["exit_code"], 0)
            out = json.loads(ret["stdout"])
            self.assertEqual(out["column_names"], ["a", "b"])
            self.assertGreaterEqual(out["data_rows"], 2)
        finally:
            Path(p).unlink(missing_ok=True)

    def test_structured_pick_json(self) -> None:
        with tempfile.NamedTemporaryFile("w", suffix=".json", delete=False, encoding="utf-8") as f:
            json.dump({"x": {"y": 42}, "arr": [1, {"z": "ok"}]}, f)
            p = f.name
        try:
            ret = run_skill_call(f"structured_pick({p!r}, pointer='x.y')", self.funcs)
            self.assertEqual(ret["exit_code"], 0)
            out = json.loads(ret["stdout"])
            self.assertEqual(out["value"], 42)
            ret2 = run_skill_call(f"structured_pick({p!r}, pointer='arr.1.z')", self.funcs)
            self.assertEqual(json.loads(ret2["stdout"])["value"], "ok")
        finally:
            Path(p).unlink(missing_ok=True)

    def test_git_snapshot_log(self) -> None:
        root = str(_root())
        ret = run_skill_call(f"git_snapshot('log', repo_path={root!r}, limit=3)", self.funcs)
        self.assertEqual(ret["exit_code"], 0)
        out = json.loads(ret["stdout"])
        self.assertIn("output", out)
        self.assertTrue(len(out["output"]) > 0)

    def test_xlsx_sample_optional(self) -> None:
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            self.skipTest("openpyxl 未安装")
        from openpyxl import Workbook

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            p = f.name
        try:
            wb = Workbook()
            ws = wb.active
            ws.append(["h1", "h2"])
            ws.append(["a", "b"])
            wb.save(p)
            wb.close()
            ret = run_skill_call(f"xlsx_sample({p!r}, max_rows=5, max_cols=4)", self.funcs)
            self.assertEqual(ret["exit_code"], 0)
            out = json.loads(ret["stdout"])
            self.assertGreaterEqual(len(out["grid"]), 1)
        finally:
            Path(p).unlink(missing_ok=True)

    def test_docx_to_text_optional(self) -> None:
        try:
            from docx import Document  # noqa: F401
        except ImportError:
            self.skipTest("python-docx 未安装")
        from docx import Document

        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
            p = f.name
        try:
            d = Document()
            d.add_paragraph("Hello docx skill")
            d.save(p)
            ret = run_skill_call(f"docx_to_text({p!r}, max_chars=10000)", self.funcs)
            self.assertEqual(ret["exit_code"], 0)
            out = json.loads(ret["stdout"])
            self.assertIn("Hello", out["text"])
        finally:
            Path(p).unlink(missing_ok=True)

    def test_run_skill_rejects_nested_call_in_args(self) -> None:
        ret = run_skill_call("date_now(__import__('os'))", self.funcs)
        self.assertEqual(ret["exit_code"], 1)
        self.assertIn("不允许的表达式节点", ret["stderr"])

    def test_run_skill_allows_required_arg_via_keyword(self) -> None:
        ret = run_skill_call("claude_job_start(prompt='分析项目')", self.funcs)
        # 关键字参数应通过 AST/manifest 参数数校验，后续失败应来自上下文约束而非“缺少位置参数”。
        self.assertNotIn("至少需要 1 个位置参数", ret["stderr"])
        self.assertNotIn("至少需要 1 个参数", ret["stderr"])
