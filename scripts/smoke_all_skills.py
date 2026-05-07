#!/usr/bin/env python3
"""对 manifest 中每个技能尝试一次最小 `run_skill` 调用，汇总 OK / SKIP / FAIL。

用法（在仓库根目录）：
  python3 scripts/smoke_all_skills.py

可选环境变量（完整表见仓库根 **`.env.example`「九·1」**；亦可写入 `.env.local`）：
  OMLXCLI_SMOKE_NETWORK=1
  OMLXCLI_SMOKE_STOCK=1
  OMLXCLI_SMOKE_PDF_PATH=/绝对路径/样例.pdf
  OMLXCLI_SMOKE_IMAGE_PATH=/绝对路径/样例.png
  OMLXCLI_SMOKE_AUDIO_PATH=/绝对路径/样例.m4a
  OMLXCLI_SMOKE_VIDEO_PATH=/绝对路径/样例.mp4
  web_search 另需 .env 中 OMLXCLI_SEARCH_GATEWAY_* 或 OMLXCLI_SEARXNG_URL（见 .env.example 第五节）。
  web_read 需 OMLXCLI_SMOKE_NETWORK=1 且 OMLXCLI_EVAL_SKIP_HTTP 不能为 1/true。
  vision_* / audio_transcribe / video_summarize 需已配置推理端点：见「十」_AICLI_API_BASE（及 KEY、MODEL），
  未 export 时冒烟为 SKIP（与仅 CLI 跑脚本、无 Web 会话一致）。

无 OMLXCLI_SMOKE_* 路径变量：`csv_tsv_summary` / `xlsx_sample` / `docx_to_text` 在 `.omlxcli/_smoke_sample.{csv,xlsx,docx}` 自造样例（跑完删）；
`structured_pick` 读 `OI_TOOL_MAP.json`；`git_snapshot` 用 `__ROOT__` 的 git log。缺 openpyxl/python-docx 或缺 OI_TOOL_MAP.json 时 SKIP。

如何传入子进程：
  - 本脚本启动时会加载仓库根目录的 **`.env` / `.env.local`**（与 `webapi` 相同规则），写入其中的
    `OMLXCLI_SMOKE_*` 可被 Python 读到，**无需**在 shell 里再 export。
  - 若在终端临时赋值，必须用 **`export VAR=value`**，否则 `python3` 子进程**继承不到**。
  - `web_read` 还需：**不要**设置 `OMLXCLI_EVAL_SKIP_HTTP=1`（与 CI 单测跳过外网一致时请 unset 或设为 0）。

退出码：若有 FAIL 则为 1；仅 SKIP/OK 为 0。
"""

from __future__ import annotations

import importlib.util
import json
import os
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]


def _root() -> str:
    return str(ROOT).replace("\\", "/")


def _expand(expr: str) -> str:
    return expr.replace("__ROOT__", _root())


def _want_network() -> bool:
    return (os.getenv("OMLXCLI_SMOKE_NETWORK") or "").strip().lower() in ("1", "true", "yes", "on")


def _want_stock_smoke() -> bool:
    """股票 skills 冒烟默认关闭，避免外网风控导致误报。"""
    return (os.getenv("OMLXCLI_SMOKE_STOCK") or "").strip().lower() in ("1", "true", "yes", "on")


def _want_http_eval() -> bool:
    """与单测一致：若显式跳过外网评测则 web_read 仍 SKIP。"""
    return (os.getenv("OMLXCLI_EVAL_SKIP_HTTP") or "").strip().lower() not in (
        "1",
        "true",
        "yes",
        "on",
    )


def _gateway() -> bool:
    return bool(
        (os.getenv("OMLXCLI_SEARCH_GATEWAY_URL") or "").strip()
        or (os.getenv("OMLXCLI_SEARXNG_URL") or "").strip()
    )


def _has_mlx_whisper() -> bool:
    """音频/视频音轨转写依赖 `mlx-whisper` Python 包。"""
    return importlib.util.find_spec("mlx_whisper") is not None


def _smoke_has_llm_endpoint() -> bool:
    """与 `.omlxcli/skills/_media._llm_endpoint` 一致：无 Web 注入时需自行 export _AICLI_API_BASE。"""
    return bool((os.getenv("_AICLI_API_BASE") or "").strip())


def main() -> int:
    sys.path.insert(0, str(ROOT))
    os.chdir(ROOT)
    from webapi.dotenv_loader import load_dotenv_files

    load_dotenv_files(ROOT)

    from webapi.skill_runner import load_skills_registry, run_skill_call

    manifest_path = ROOT / ".omlxcli" / "skills" / "manifests" / "skills.json"
    data = json.loads(manifest_path.read_text(encoding="utf-8"))
    # note_list / note_load 可在 note_save 之前（空笔记与 list 均合法）
    skill_names = sorted((data.get("skills") or {}).keys())
    funcs, _ = load_skills_registry()
    if not funcs:
        print("FAIL: skills 未加载", file=sys.stderr)
        return 1

    rows: list[tuple[str, str, str, str]] = []
    cleanup_files: list[Path] = []
    note = "smoke_all_skills_note"

    def run(name: str, expr: str) -> tuple[str, str, str]:
        ret = run_skill_call(_expand(expr), funcs)
        code = int(ret.get("exit_code", 1))
        if code == 0:
            return name, "OK", (ret.get("stdout") or "")[:120].replace("\n", " ")
        return name, "FAIL", (ret.get("stderr") or "")[:200]

    for name in skill_names:
        if name not in funcs:
            rows.append((name, "FAIL", f"manifest 有但 builtins 未注册: {name}", ""))
            continue

        skip_reason = ""
        expr = ""

        if name == "date_now":
            expr = "date_now()"
        elif name == "note_save":
            expr = f"note_save('{note}', 'smoke line')"
        elif name == "note_load":
            expr = f"note_load('{note}')"
        elif name == "note_list":
            expr = "note_list()"
        elif name == "files_read_chunk":
            expr = "files_read_chunk('__ROOT__/README.md', start=0, lines=2)"
        elif name == "files_search":
            expr = "files_search('SessionStore', path='__ROOT__/webapi', kind='content', max_results=1)"
        elif name.startswith("claude_job_"):
            skip_reason = (
                "需 Web 会话与 Claude Job 服务；冒烟不测（见 docs/CLAUDE_CODE_JOB_SPEC.md）"
            )
        elif name == "repo_grep":
            expr = "repo_grep('def', path='__ROOT__/webapi', max_matches=2)"
        elif name == "csv_tsv_summary":
            p = ROOT / ".omlxcli" / "_smoke_sample.csv"
            p.parent.mkdir(parents=True, exist_ok=True)
            p.write_text("h1,h2\n3,4\n", encoding="utf-8")
            cleanup_files.append(p)
            expr = f"csv_tsv_summary({json.dumps(str(p))})"
        elif name == "structured_pick":
            jpath = ROOT / "OI_TOOL_MAP.json"
            if not jpath.is_file():
                skip_reason = "缺少 OI_TOOL_MAP.json（请先 gen_oi_tool_map --write）"
            else:
                expr = f"structured_pick({json.dumps(str(jpath))}, pointer='scope', format='json')"
        elif name == "git_snapshot":
            expr = "git_snapshot('log', repo_path='__ROOT__', limit=5)"
        elif name == "xlsx_sample":
            try:
                from openpyxl import Workbook
            except ImportError:
                skip_reason = "缺少 openpyxl（见 requirements.txt）"
            else:
                px = ROOT / ".omlxcli" / "_smoke_sample.xlsx"
                wb = Workbook()
                ws = wb.active
                ws.append(["c", "d"])
                ws.append(["1", "2"])
                wb.save(px)
                wb.close()
                cleanup_files.append(px)
                expr = f"xlsx_sample({json.dumps(str(px))}, max_rows=6, max_cols=4)"
        elif name == "docx_to_text":
            try:
                from docx import Document
            except ImportError:
                skip_reason = "缺少 python-docx（见 requirements.txt）"
            else:
                px = ROOT / ".omlxcli" / "_smoke_sample.docx"
                d = Document()
                d.add_paragraph("smoke docx line")
                d.save(px)
                cleanup_files.append(px)
                expr = f"docx_to_text({json.dumps(str(px))}, max_chars=8000)"
        elif name in ("web_search",):
            if not _want_network() or not _gateway():
                skip_reason = "需外网 + 网关：设置 OMLXCLI_SMOKE_NETWORK=1 且配置 SEARCH_GATEWAY 或 SEARXNG"
            else:
                expr = "web_search('Python', max_results=1, fetch_content=False)"
        elif name in ("web_read",):
            if not _want_network() or not _want_http_eval():
                skip_reason = (
                    "需 OMLXCLI_SMOKE_NETWORK=1，且 OMLXCLI_EVAL_SKIP_HTTP 不能为 1/true"
                    "（请 unset 或设为 0；样例路径可放 .env.local）"
                )
            else:
                expr = "web_read('https://example.com', max_chars=400)"
        elif name in ("weather_now", "weather_forecast"):
            if not _want_network():
                skip_reason = "需外网：OMLXCLI_SMOKE_NETWORK=1"
            elif name == "weather_now":
                expr = "weather_now('Beijing')"
            else:
                expr = "weather_forecast('Beijing', days=2)"
        elif name.startswith("stock_"):
            if not _want_network() or not _want_stock_smoke():
                skip_reason = "股票技能需 OMLXCLI_SMOKE_NETWORK=1 且 OMLXCLI_SMOKE_STOCK=1（默认关闭避免风控误报）"
            elif name == "stock_search":
                expr = "stock_search('贵州茅台', limit=3)"
            elif name == "stock_quote":
                expr = "stock_quote('600519', instrument='stock')"
            elif name == "stock_hot_list":
                expr = "stock_hot_list(limit=5)"
            elif name == "stock_unusual":
                expr = "stock_unusual(limit=10)"
            elif name == "stock_brief":
                expr = "stock_brief('600519')"
            elif name == "stock_kline":
                expr = "stock_kline('600519', period='day', count=30, adjust='qfq')"
            elif name == "stock_history_trades":
                expr = "stock_history_trades('600519', count=30)"
            elif name == "stock_kline_summary":
                expr = "stock_kline_summary('600519', period='day', bars=30)"
            else:
                skip_reason = f"未配置表达式: {name}"
        elif name.startswith("pdf_"):
            pdf = (os.getenv("OMLXCLI_SMOKE_PDF_PATH") or "").strip()
            if not pdf or not Path(pdf).is_file():
                skip_reason = "无 PDF：设置 OMLXCLI_SMOKE_PDF_PATH=/绝对路径/样例.pdf"
            elif name == "pdf_meta":
                expr = f"pdf_meta({json.dumps(pdf)})"
            elif name == "pdf_read":
                expr = f"pdf_read({json.dumps(pdf)}, pages=(1, 1), ocr='off')"
            elif name == "pdf_to_text":
                expr = f"pdf_to_text({json.dumps(pdf)}, pages=(1, 1))"
            elif name == "pdf_ocr":
                expr = f"pdf_ocr({json.dumps(pdf)}, pages=(1, 1))"
            elif name == "pdf_search":
                expr = f"pdf_search({json.dumps(pdf)}, 'the')"
            else:
                skip_reason = f"未配置表达式: {name}"
        elif name.startswith("vision_"):
            img = (os.getenv("OMLXCLI_SMOKE_IMAGE_PATH") or "").strip()
            if not img or not Path(img).is_file():
                skip_reason = "无图片：设置 OMLXCLI_SMOKE_IMAGE_PATH=/绝对路径/样例.png"
            elif not _smoke_has_llm_endpoint():
                skip_reason = "缺推理端点：export _AICLI_API_BASE（及 _AICLI_API_KEY、_AICLI_LLM_MODEL），见 .env.example「十」"
            elif name == "vision_describe":
                expr = f"vision_describe({json.dumps(img)}, '用一句话描述画面')"
            else:
                expr = f"vision_compare([{json.dumps(img)}], '图中主色是什么？')"
        elif name.startswith("audio_"):
            aud = (os.getenv("OMLXCLI_SMOKE_AUDIO_PATH") or "").strip()
            if not aud or not Path(aud).is_file():
                skip_reason = "无音频：设置 OMLXCLI_SMOKE_AUDIO_PATH=/绝对路径/样例.m4a"
            elif not _has_mlx_whisper():
                skip_reason = (
                    "当前解释器未检测到 mlx-whisper：Apple Silicon 上请 `./bootstrap.sh` 后用 "
                    "`.venv/bin/python` 跑本脚本；或 `pip install mlx-whisper`（需与 Web 同一 Python）"
                )
            elif name == "audio_transcribe_only":
                expr = f"audio_transcribe_only({json.dumps(aud)})"
            elif not _smoke_has_llm_endpoint():
                skip_reason = "缺推理端点：audio_transcribe 需 LLM 总结；export _AICLI_* 见 .env.example「十」"
            else:
                expr = f"audio_transcribe({json.dumps(aud)}, '用一句话总结', language='zh')"
        elif name == "video_summarize":
            vid = (os.getenv("OMLXCLI_SMOKE_VIDEO_PATH") or "").strip()
            if not vid or not Path(vid).is_file():
                skip_reason = "无视频：设置 OMLXCLI_SMOKE_VIDEO_PATH=/绝对路径/样例.mp4"
            elif not _smoke_has_llm_endpoint():
                skip_reason = "缺推理端点：export _AICLI_API_BASE 等，见 .env.example「十」"
            else:
                # 音轨缺 mlx-whisper 时 skill 内部会跳过转写，仍可做画面分析
                expr = f"video_summarize({json.dumps(vid)}, frames=2, prompt='一句话概括', language='zh')"

        if skip_reason:
            rows.append((name, "SKIP", skip_reason, ""))
            continue
        if not expr:
            rows.append((name, "SKIP", "未实现该技能的 smoke 表达式", ""))
            continue
        n, st, detail = run(name, expr)
        rows.append((n, st, detail, expr))

    for fp in cleanup_files:
        try:
            fp.unlink(missing_ok=True)
        except OSError:
            pass

    # 清理笔记（见 `.omlxcli/skills/notes.py` → `.omlxcli/notes/`）
    try:
        nd = ROOT / ".omlxcli" / "notes"
        for p in nd.glob(f"{note}*"):
            if p.is_file():
                p.unlink()
    except OSError:
        pass

    fail = 0
    for n, st, detail, _ex in rows:
        line = f"{n:22} {st:6}"
        if detail:
            line += f"  {detail}"
        print(line)
        if st == "FAIL":
            fail += 1

    print("---")
    print(f"合计: {len(rows)}  FAIL={fail}")
    return 1 if fail else 0


if __name__ == "__main__":
    raise SystemExit(main())
